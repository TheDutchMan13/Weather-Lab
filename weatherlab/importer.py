"""Import past readings from a file the user uploads. Nothing else is read from the computer.

Supported files: CSV and Excel (.xlsx). Each file needs a time column named time, timestamp,
datetime, or date. Other columns are matched to readings by name (see SYNONYMS). The user can
change any match on the import page. Values are stored in the app's own units (imperial), so
metric files are converted when the user says so.
"""

import csv
import io
import json
import re
from datetime import datetime
from zoneinfo import ZoneInfo

from weatherlab import daily, parse

TIME_COLUMNS = ("time", "timestamp", "datetime", "date")
READING_KEYS = ("tempf", "humidity", "feelslikef", "dewpointf", "windspeedmph", "windgustmph", "winddir",
                "dailyrainin", "rainratein", "solarradiation", "uv", "tempinf", "humidityin",
                "baromrelin", "baromabsin")
# Common column names for the same reading. Names are compared after _norm() removes spaces,
# punctuation, and any unit in brackets, so "Dew Point" and "Dew point (°F)" both match.
SYNONYMS = {
    "temperature": "tempf", "outdoortemperature": "tempf", "temp": "tempf",
    "dewpoint": "dewpointf",
    "feelslike": "feelslikef",
    "humidity": "humidity", "outdoorhumidity": "humidity", "relativehumidity": "humidity",
    "windspeed": "windspeedmph", "wind": "windspeedmph",
    "windgust": "windgustmph", "gust": "windgustmph",
    "winddirection": "winddir", "winddir": "winddir",
    "preciprate": "rainratein", "rainrate": "rainratein",
    "dailyrain": "dailyrainin", "raintoday": "dailyrainin",
    "solar": "solarradiation", "solarradiation": "solarradiation",
    "uv": "uv", "uvindex": "uv",
    "pressure": "baromrelin", "barometer": "baromrelin", "sealevelpressure": "baromrelin",
    "indoortemperature": "tempinf", "indoorhumidity": "humidityin",
}
# Export column order. Only columns that have values are written.
EXPORT_KEYS = READING_KEYS
MAX_BYTES = 20 * 1024 * 1024
SOURCE = "import"

# Metric value -> the imperial value the app stores. Keys not listed are already unitless or the same.
METRIC_TO_IMPERIAL = {
    "tempf": lambda c: c * 9 / 5 + 32,
    "dewpointf": lambda c: c * 9 / 5 + 32,
    "feelslikef": lambda c: c * 9 / 5 + 32,
    "tempinf": lambda c: c * 9 / 5 + 32,
    "windspeedmph": lambda v: v * 0.621371,
    "windgustmph": lambda v: v * 0.621371,
    "baromrelin": lambda hpa: hpa * 0.0295300,
    "baromabsin": lambda hpa: hpa * 0.0295300,
    "rainratein": lambda mm: mm * 0.0393701,
    "dailyrainin": lambda mm: mm * 0.0393701,
}


class ImportProblem(Exception):
    """The file can't be imported. The message is shown to the user as it is."""


def _norm(text: str) -> str:
    text = re.sub(r"\(.*?\)", "", str(text))
    return re.sub(r"[^a-z0-9]", "", text.lower())


def guess_key(header: str):
    n = _norm(header)
    if n in READING_KEYS:
        return n
    return SYNONYMS.get(n)


def guess_time_column(headers: list[str]):
    for h in headers:
        if _norm(h) in TIME_COLUMNS:
            return h
    return None


def guess_mapping(headers: list[str], time_col) -> dict[str, str]:
    """Column name -> reading key, for every column that matches and isn't used twice."""
    out, used = {}, set()
    for h in headers:
        if h == time_col:
            continue
        key = guess_key(h)
        if key and key not in used:
            out[h] = key
            used.add(key)
    return out


def _num(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def parse_time(value, tz: ZoneInfo) -> int:
    if isinstance(value, datetime):
        when = value if value.tzinfo else value.replace(tzinfo=tz)
        return int(when.timestamp())
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip()
    try:
        return int(float(text))
    except ValueError:
        pass
    try:
        when = datetime.fromisoformat(text)
    except ValueError:
        raise ImportProblem(f"Can't read the time '{text}'. Use Unix seconds or a date like 2026-10-20 13:45.")
    if when.tzinfo is None:
        when = when.replace(tzinfo=tz)
    return int(when.timestamp())


def _read_csv(raw: bytes):
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ImportProblem("Save the file as a UTF-8 CSV and try again.")
    reader = csv.DictReader(io.StringIO(text))
    headers = [h.strip() for h in (reader.fieldnames or []) if h and h.strip()]
    rows = [{(k or "").strip(): v for k, v in rec.items() if k} for rec in reader]
    return headers, rows


def _read_xlsx(raw: bytes):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportProblem("Excel import needs the openpyxl package. Run: pip install -r requirements.txt")
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        values = workbook.worksheets[0].iter_rows(values_only=True)
        head = next(values, None)
        if head is None:
            raise ImportProblem("The first sheet is empty.")
        names = [str(h).strip() if h is not None and str(h).strip() else f"Column {i + 1}"
                 for i, h in enumerate(head)]
        rows = [dict(zip(names, vals)) for vals in values if any(v is not None for v in vals)]
    except ImportProblem:
        raise
    except Exception as exc:
        raise ImportProblem("Couldn't open the Excel file. Save it as .xlsx and try again.") from exc
    return names, rows


def read_file(raw: bytes, filename: str = ""):
    """(headers, rows) for a CSV or .xlsx file. Each row is a dict keyed by header."""
    if len(raw) > MAX_BYTES:
        raise ImportProblem("The file is larger than 20 MB. Split it into smaller files.")
    if filename.lower().endswith(".xlsx"):
        return _read_xlsx(raw)
    return _read_csv(raw)


def build_readings(rows: list[dict], time_col: str, mapping: dict, tz: ZoneInfo,
                   units: str = "imperial") -> list[tuple[int, dict]]:
    """Turn rows into (ts, reading) pairs. Metric values are converted to imperial when units is 'metric'."""
    conv = METRIC_TO_IMPERIAL if units == "metric" else {}
    out = []
    for line_no, rec in enumerate(rows, start=2):
        raw_time = rec.get(time_col)
        if raw_time is None or str(raw_time).strip() == "":
            continue
        try:
            ts = parse_time(raw_time, tz)
        except ImportProblem as exc:
            raise ImportProblem(f"Line {line_no}: {exc}")
        reading = {}
        for header, key in mapping.items():
            value = _num(rec.get(header))
            if value is None:
                continue
            reading[key] = conv.get(key, lambda v: v)(value)
        reading = parse.clean(reading)
        if reading:
            out.append((ts, reading))
    if not out:
        raise ImportProblem("No readings found. Check the time column and the matched columns have values.")
    return out


def read_csv(raw: bytes, tz: ZoneInfo):
    """Imports a CSV with the automatic column matching. Returns (readings, ignored column names)."""
    headers, rows = read_file(raw, "upload.csv")
    time_col = guess_time_column(headers)
    if time_col is None:
        raise ImportProblem("No time column found. Name one column time, timestamp, datetime, or date.")
    mapping = guess_mapping(headers, time_col)
    readings = build_readings(rows, time_col, mapping, tz)
    ignored = [h for h in headers if h != time_col and h not in mapping]
    return readings, ignored


def export_csv(rows: list[tuple[int, dict]], tz: ZoneInfo) -> bytes:
    """A CSV the CSV Import page reads back. Times are local, in the timezone from Settings."""
    keys = [k for k in EXPORT_KEYS if any(k in r for _, r in rows)]
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["time", *keys])
    for ts, reading in rows:
        writer.writerow([datetime.fromtimestamp(ts, tz).strftime("%Y-%m-%d %H:%M:%S"),
                         *[reading.get(k, "") for k in keys]])
    return buf.getvalue().encode("utf-8")


def store(db, rows: list[tuple[int, dict]], lat: float, lon: float, elev_m: float, tz: ZoneInfo) -> dict:
    """Save the readings and rebuild the daily summaries for the days they cover."""
    before = db.total_changes
    with db:
        db.executemany("INSERT OR IGNORE INTO readings (ts, data, source) VALUES (?, ?, ?)",
                       [(ts, json.dumps(r), SOURCE) for ts, r in rows])
    inserted = db.total_changes - before
    days = sorted({datetime.fromtimestamp(ts, tz).date() for ts, _ in rows})
    if days:
        daily.rollup_days(db, days, lat, lon, elev_m, tz, datetime.now(tz).date())
    return {"found": len(rows), "inserted": inserted, "skipped": len(rows) - inserted, "days": len(days)}
